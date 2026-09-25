"""Leitura nativa de arquivos em Python + preparação de imagens para o VLM.

Se um formato não puder ser lido aqui, `UnreadableFile` traz uma sugestão de
comando de terminal; a IA então usa a ferramenta `run_terminal_command`.
"""

from __future__ import annotations

import base64
import csv
import email
import email.policy
import html
import io
import json
import mimetypes
import re
import sys
import tarfile
import zipfile
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import Path

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".jfif", ".ico"}

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json", ".jsonl", ".ndjson", ".xml", ".yaml",
    ".yml", ".toml", ".ini", ".cfg", ".conf", ".env", ".properties", ".py", ".pyw", ".js", ".mjs", ".cjs", ".ts",
    ".tsx", ".jsx", ".java", ".kt", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs", ".go", ".rs", ".rb", ".php", ".swift",
    ".sql", ".sh", ".bash", ".zsh", ".bat", ".cmd", ".ps1", ".psm1", ".css", ".scss", ".less", ".vue", ".svelte",
    ".tex", ".bib", ".srt", ".vtt", ".r", ".m", ".lua", ".pl", ".dart", ".scala", ".gradle", ".dockerfile",
    ".gitignore", ".svg", ".geojson", ".gpx", ".kml", ".nfo", ".asm", ".vb", ".vbs", ".f90", ".jl",
}

MAGIC = [
    (b"%PDF", "documento PDF"), (b"PK\x03\x04", "arquivo ZIP/Office Open XML"), (b"\x89PNG", "imagem PNG"),
    (b"\xff\xd8\xff", "imagem JPEG"), (b"GIF8", "imagem GIF"), (b"7z\xbc\xaf\x27\x1c", "arquivo 7-Zip"),
    (b"Rar!", "arquivo RAR"), (b"MZ", "executável/biblioteca Windows (PE)"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "documento OLE2 (Office 97-2003: .doc/.xls/.ppt/.msg)"),
    (b"ID3", "áudio MP3"), (b"OggS", "áudio/vídeo Ogg"), (b"fLaC", "áudio FLAC"), (b"\x1f\x8b", "arquivo GZIP"),
    (b"BZh", "arquivo BZIP2"), (b"\xfd7zXZ", "arquivo XZ"), (b"SQLite format 3", "banco SQLite"),
    (b"{\\rtf", "documento RTF"), (b"\x7fELF", "executável ELF"),
]


class UnreadableFile(Exception):
    def __init__(self, reason: str, hint: str = ""):
        super().__init__(reason)
        self.reason = reason
        self.hint = hint


def human_size(n: int) -> str:
    size = float(n)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{size:.0f} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{n} B"


def sniff_type(path: Path) -> str:
    try:
        with path.open("rb") as f:
            head = f.read(16)
    except OSError:
        return "desconhecido"
    if head[4:8] == b"ftyp":
        return "vídeo/áudio MP4/MOV (ou imagem HEIC)"
    if head[:4] == b"RIFF":
        return {b"WEBP": "imagem WEBP", b"WAVE": "áudio WAV", b"AVI ": "vídeo AVI"}.get(head[8:12], "arquivo RIFF")
    for magic, desc in MAGIC:
        if head.startswith(magic):
            return desc
    mime, _ = mimetypes.guess_type(path.name)
    return mime or "binário desconhecido"


def is_image(path: Path) -> bool:
    return path.suffix.lower() in IMAGE_EXTS


# --------------------------------------------------------------------------- texto
def _decode_bytes(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes  # type: ignore

        best = from_bytes(data[:2_000_000]).best()
        # Para alfabetos latinos, o palpite em textos curtos oscila entre cp1250/cp1252/
        # latin-2...; no Windows em português o correto é quase sempre o cp1252.
        if best is not None and best.encoding.replace("-", "_") not in _LATIN_SINGLE_BYTE:
            return data.decode(best.encoding, errors="replace")
    except Exception:  # noqa: BLE001
        pass
    return data.decode("cp1252", errors="replace")


_LATIN_SINGLE_BYTE = {
    "cp1250", "cp1252", "cp1254", "cp1257", "cp850", "cp858", "latin_1", "iso8859_1", "iso8859_2",
    "iso8859_3", "iso8859_4", "iso8859_9", "iso8859_10", "iso8859_13", "iso8859_14", "iso8859_15",
    "iso8859_16", "mac_roman", "mac_latin2", "ascii",
}


def _looks_textual(data: bytes) -> bool:
    sample = data[:8192]
    if not sample:
        return True
    if b"\x00" in sample and not sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return False
    control = sum(1 for b in sample if b < 9 or (13 < b < 32))
    return control / len(sample) < 0.05


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "head", "svg"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table",
             "title", "blockquote", "pre"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def html_to_text(markup: str) -> str:
    parser = _HTMLText()
    try:
        parser.feed(markup)
        parser.close()
    except Exception:  # noqa: BLE001
        return re.sub(r"<[^>]+>", " ", markup)
    text = "".join(parser.parts)
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _xml_to_text(markup: str) -> str:
    markup = re.sub(r"</(text:p|text:h|w:p|a:p)>", "\n", markup)
    text = html.unescape(re.sub(r"<[^>]+>", "", markup))
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


# --------------------------------------------------------------------------- formatos
def _pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as e:
        raise UnreadableFile("biblioteca pypdf não instalada", terminal_hint(path)) from e
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:  # noqa: BLE001
                raise UnreadableFile("PDF protegido por senha", terminal_hint(path)) from e
        pages = []
        for i, page in enumerate(reader.pages, 1):
            try:
                txt = page.extract_text() or ""
            except Exception:  # noqa: BLE001
                txt = ""
            pages.append(f"--- Página {i} ---\n{txt.strip()}")
    except UnreadableFile:
        raise
    except Exception as e:  # noqa: BLE001
        raise UnreadableFile(f"falha ao abrir o PDF ({e})", terminal_hint(path)) from e
    text = "\n\n".join(pages)
    if len(re.sub(r"--- Página \d+ ---|\s", "", text)) < 20 * max(1, len(pages)) // 4:
        raise UnreadableFile("PDF sem camada de texto (provavelmente escaneado; precisa de OCR)",
                             terminal_hint(path, ocr=True))
    return text


def _docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError:
        with zipfile.ZipFile(path) as z:
            return _xml_to_text(z.read("word/document.xml").decode("utf-8", errors="replace"))
    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t_i, table in enumerate(doc.tables, 1):
        parts.append(f"\n[Tabela {t_i}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts).strip()


def _xlsx(path: Path, max_rows: int = 2000) -> str:
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            out.append(f"=== Planilha: {ws.title} ===")
            for r_i, row in enumerate(ws.iter_rows(values_only=True)):
                if r_i >= max_rows:
                    out.append(f"... (linhas além de {max_rows} omitidas)")
                    break
                if row is None or all(v is None for v in row):
                    continue
                out.append("\t".join("" if v is None else str(v) for v in row))
    finally:
        wb.close()
    return "\n".join(out)


def _pptx(path: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        with zipfile.ZipFile(path) as z:
            names = sorted(n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))
            return "\n\n".join(_xml_to_text(z.read(n).decode("utf-8", errors="replace")) for n in names)
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Notas: " + slide.notes_slide.notes_text_frame.text.strip())
    return "\n".join(out)


def _odf(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return _xml_to_text(z.read("content.xml").decode("utf-8", errors="replace"))


def _epub(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm"))]
        return "\n\n".join(html_to_text(z.read(n).decode("utf-8", errors="replace")) for n in sorted(names))


def _rtf(path: Path) -> str:
    raw = path.read_bytes().decode("latin-1", errors="replace")
    raw = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1252", "replace"), raw)
    raw = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), raw)
    raw = re.sub(r"\\(par|line)\b ?", "\n", raw)
    raw = re.sub(r"\{\\\*[^{}]*\}", "", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", raw)
    return re.sub(r"[{}]", "", raw).strip()


def _eml(path: Path) -> str:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    head = "\n".join(f"{k}: {msg[k]}" for k in ("From", "To", "Cc", "Date", "Subject") if msg[k])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        text = body.get_content()
        if body.get_content_type() == "text/html":
            text = html_to_text(text)
    atts = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
    extra = f"\n\nAnexos: {', '.join(atts)}" if atts else ""
    return f"{head}\n\n{text}{extra}".strip()


def _ipynb(path: Path) -> str:
    nb = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    out = []
    for i, cell in enumerate(nb.get("cells", []), 1):
        src = "".join(cell.get("source", []))
        out.append(f"--- Célula {i} [{cell.get('cell_type')}] ---\n{src}")
    return "\n\n".join(out)


def _archive_listing(path: Path) -> str:
    lines = []
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            for info in z.infolist()[:500]:
                lines.append(f"{info.file_size:>12,}  {info.filename}")
        return "Conteúdo do arquivo compactado (ZIP):\n" + "\n".join(lines)
    with tarfile.open(path) as t:
        for m in t.getmembers()[:500]:
            lines.append(f"{m.size:>12,}  {m.name}")
    return "Conteúdo do arquivo compactado (TAR):\n" + "\n".join(lines)


def _csv(path: Path, max_rows: int = 5000) -> str:
    text = _decode_bytes(path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        return text
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= max_rows:
            rows.append(f"... (linhas além de {max_rows} omitidas)")
            break
        rows.append("\t".join(row))
    return "\n".join(rows)


_HANDLERS = {
    ".pdf": _pdf, ".docx": _docx, ".docm": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pptx": _pptx,
    ".odt": _odf, ".ods": _odf, ".odp": _odf, ".epub": _epub, ".rtf": _rtf, ".eml": _eml, ".ipynb": _ipynb,
    ".zip": _archive_listing, ".tar": _archive_listing, ".tgz": _archive_listing, ".gz": _archive_listing,
    ".bz2": _archive_listing, ".xz": _archive_listing, ".csv": _csv,
}


def extract_text(path: str | Path) -> str:
    """Texto do arquivo, ou UnreadableFile com uma sugestão de comando de terminal."""
    path = Path(path)
    if not path.exists():
        raise UnreadableFile(f"arquivo não encontrado: {path}")
    if path.is_dir():
        raise UnreadableFile("o caminho é uma pasta", "Use run_terminal_command para listar (ex.: dir / ls).")
    st = path.stat()
    return _extract_cached(str(path), st.st_mtime, st.st_size)


@lru_cache(maxsize=32)
def _extract_cached(path_str: str, _mtime: float, size: int) -> str:
    path = Path(path_str)
    ext = path.suffix.lower()
    if ext in (".html", ".htm", ".xhtml", ".mhtml"):
        return html_to_text(_decode_bytes(path.read_bytes()))
    handler = _HANDLERS.get(ext)
    if handler is not None:
        try:
            return handler(path)
        except UnreadableFile:
            raise
        except Exception as e:  # noqa: BLE001
            raise UnreadableFile(f"falha ao ler {ext} ({type(e).__name__}: {e})", terminal_hint(path)) from e
    if size > 200 * 2**20:
        raise UnreadableFile("arquivo muito grande para leitura direta", terminal_hint(path))
    data = path.read_bytes()
    if ext in TEXT_EXTS or _looks_textual(data):
        return _decode_bytes(data)
    raise UnreadableFile(f"formato binário não suportado nativamente ({sniff_type(path)})", terminal_hint(path))


def terminal_hint(path: Path, ocr: bool = False) -> str:
    """Sugestões de comandos para extrair o conteúdo via terminal."""
    p, ext = str(path), path.suffix.lower()
    win = sys.platform == "win32"
    tips: list[str] = []
    if ext == ".pdf":
        if ocr:
            tips.append(f'OCR: tesseract "{p}" - -l por  (converta páginas em imagem antes, ex.: pdftoppm -png)')
            tips.append(f'ou: ocrmypdf -l por "{p}" saida.pdf  e depois  pdftotext -layout saida.pdf -')
        tips.append(f'pdftotext -layout "{p}" -   (Poppler)')
    elif ext in (".doc", ".xls", ".ppt", ".rtf", ".odt", ".wpd", ".pages", ".xlsb"):
        tips.append(f'soffice --headless --convert-to txt:Text --outdir "{path.parent}" "{p}"   (LibreOffice)')
        if win and ext == ".doc":
            tips.append("PowerShell + Word: $w=New-Object -ComObject Word.Application; "
                        f"$d=$w.Documents.Open('{p}'); $d.Content.Text; $d.Close(); $w.Quit()")
        if win and ext in (".xls", ".xlsb"):
            tips.append("PowerShell + Excel: $x=New-Object -ComObject Excel.Application; "
                        f"$b=$x.Workbooks.Open('{p}'); $b.SaveAs('{path.with_suffix('.csv')}',6); $b.Close(); $x.Quit()")
    elif ext == ".msg":
        tips.append("PowerShell + Outlook: $o=New-Object -ComObject Outlook.Application; "
                    f"$m=$o.Session.OpenSharedItem('{p}'); $m.Subject; $m.Body")
    elif ext in (".7z", ".rar", ".cab", ".iso"):
        tips.append(f'tar -tf "{p}"   (lista o conteúdo; o tar do Windows 10+ lê vários formatos)')
        tips.append(f'7z l "{p}"   (7-Zip)')
    elif ext in (".mp3", ".wav", ".flac", ".ogg", ".m4a", ".mp4", ".mkv", ".mov", ".avi", ".webm"):
        tips.append(f'ffprobe -hide_banner "{p}"   (metadados; requer FFmpeg)')
    elif ext in (".exe", ".dll", ".msi", ".sys"):
        tips.append(f'(Get-Item "{p}").VersionInfo | Format-List' if win else f'file "{p}"')
    elif ext in (".heic", ".heif", ".avif", ".psd", ".raw", ".cr2", ".nef"):
        tips.append(f'magick "{p}" convertido.png   (ImageMagick) — depois peça ao usuário para anexar o PNG')
    if not tips or not ocr:
        tips.append(f'Format-Hex -Path "{p}" | Select-Object -First 30' if win else f'xxd "{p}" | head -30')
        if not win:
            tips.append(f'strings -n 6 "{p}" | head -200')
    return "Sugestões: " + " || ".join(tips)


# --------------------------------------------------------------------------- imagens
def image_to_data_uri(path: str | Path, max_side: int = 1280) -> str:
    """Reduz a imagem (menos tokens e menos VRAM no encoder) e devolve um data URI."""
    path = Path(path)
    try:
        from PIL import Image, ImageOps  # type: ignore

        with Image.open(path) as im:
            im = ImageOps.exif_transpose(im)
            if getattr(im, "n_frames", 1) > 1:
                im.seek(0)
            if im.mode not in ("RGB", "L"):
                background = Image.new("RGB", im.size, (255, 255, 255))
                rgba = im.convert("RGBA")
                background.paste(rgba, mask=rgba.split()[-1])
                im = background
            elif im.mode == "L":
                im = im.convert("RGB")
            if max(im.size) > max_side > 0:
                im.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
            buf = io.BytesIO()
            im.save(buf, format="JPEG", quality=92)
            return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode("ascii")
    except ImportError:
        mime = mimetypes.guess_type(path.name)[0] or "image/png"
        return f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode("ascii")
